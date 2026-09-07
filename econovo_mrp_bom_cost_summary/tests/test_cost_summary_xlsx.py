# -*- coding: utf-8 -*-
"""Verify the Excel export mirrors the interactive UI's cost visibility
(omit cost columns for users without "Show Product Cost") instead of denying
the whole file, which is what the endpoint did before.

Uses real BOM data from ``get_html`` (same fixture pattern as
``test_cost_summary.py``) and calls the private sheet builders directly,
sidestepping the HTTP/session layer entirely: the change under test is which
columns the builders emit, not the controller's request plumbing.
"""
from openpyxl import Workbook

from odoo.tests import tagged
from odoo.tests.common import TransactionCase

from odoo.addons.econovo_mrp_bom_cost_summary.controllers.bom_cost_summary_xlsx import (
    _build_detail_sheet,
    _build_summary_sheet,
    _build_tree_sheet,
)


@tagged("post_install", "-at_install")
class TestBomCostSummaryXlsxCostGating(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.report = cls.env["report.mrp.report_bom_structure"]
        cls.summary_model = cls.env[
            "report.econovo_mrp_bom_cost_summary.report_cost_summary"
        ]
        cls.uom_unit = cls.env.ref("uom.product_uom_unit")
        cls.group_show_cost = cls.env.ref(
            "hide_product_price_cost."
            "hide_product_price_cost_group_user_show_product_cost"
        )
        cls.env.user.groups_id |= cls.group_show_cost

        cls.comp = cls.env["product.product"].create({
            "name": "XLSX Gating Component",
            "type": "product",
            "standard_price": 100.0,
            "uom_id": cls.uom_unit.id,
            "uom_po_id": cls.uom_unit.id,
        })
        cls.finished = cls.env["product.product"].create({
            "name": "XLSX Gating Finished",
            "type": "product",
            "standard_price": 500.0,
            "uom_id": cls.uom_unit.id,
            "uom_po_id": cls.uom_unit.id,
        })
        cls.bom = cls.env["mrp.bom"].create({
            "product_tmpl_id": cls.finished.product_tmpl_id.id,
            "product_qty": 1.0,
            "product_uom_id": cls.uom_unit.id,
            "type": "normal",
            "bom_line_ids": [(0, 0, {
                "product_id": cls.comp.id,
                "product_qty": 2.0,
            })],
        })

    def _raw_and_summary(self):
        raw = self.report.get_html(bom_id=self.bom.id, searchQty=1)
        return raw["lines"], raw["lines"]["cost_summary"]

    def _headers(self, sheet):
        # The "BOM Tree" sheet writes 2 info rows + 1 blank separator before
        # the actual column-header row, so scan the whole sheet rather than
        # assume the headers sit on row 1 (true only for some builders).
        return [
            cell.value for row in sheet.iter_rows() for cell in row
            if isinstance(cell.value, str)
        ]

    def _build_all_sheets(self, bom_lines, cost_summary, show_costs):
        wb = Workbook()
        tree_ws = wb.active
        _build_tree_sheet(
            tree_ws, bom_lines, "ARS", "", 0,
            show_costs, True, True, self.bom.display_name, 1,
        )
        summary_ws = wb.create_sheet("Cost Summary")
        _build_summary_sheet(
            summary_ws, cost_summary, "ARS", "",
            show_costs, True, True, self.bom.display_name, 1,
        )
        detail_ws = wb.create_sheet("Components Detail")
        _build_detail_sheet(detail_ws, cost_summary, "ARS", "", show_costs, True)
        return tree_ws, summary_ws, detail_ws

    def test_with_group_includes_cost_columns(self):
        bom_lines, cost_summary = self._raw_and_summary()
        self.assertTrue(cost_summary["show_costs"])
        tree_ws, summary_ws, detail_ws = self._build_all_sheets(
            bom_lines, cost_summary, show_costs=True,
        )
        for sheet in (tree_ws, summary_ws, detail_ws):
            self.assertTrue(
                any("Cost (" in (h or "") for h in self._headers(sheet)),
                "%s must expose cost columns when show_costs is True"
                % sheet.title,
            )

    def test_without_group_omits_cost_columns(self):
        self.env.user.groups_id -= self.group_show_cost
        bom_lines, cost_summary = self._raw_and_summary()
        self.assertFalse(cost_summary["show_costs"])
        # The controller ANDs the client's display toggle with the group, so
        # a user without the group always resolves to show_costs=False.
        tree_ws, summary_ws, detail_ws = self._build_all_sheets(
            bom_lines, cost_summary, show_costs=False,
        )
        for sheet in (tree_ws, summary_ws, detail_ws):
            self.assertFalse(
                any("Cost (" in (h or "") for h in self._headers(sheet)),
                "%s must not expose any cost column without the group"
                % sheet.title,
            )
            self.assertTrue(
                self._headers(sheet),
                "%s must still expose its structural columns" % sheet.title,
            )
